"""上传文件解析。txt/md/csv 零依赖；docx/pdf 依赖可选库，缺失时给出明确错误而不是静默失败。"""
from __future__ import annotations

import csv
import io
from pathlib import PurePosixPath

MAX_BYTES = 8 * 1024 * 1024
SUPPORTED = (".txt", ".md", ".markdown", ".csv", ".docx", ".pdf")


class ImportError_(ValueError):
    """导入失败（文件名、类型、编码、内容为空、缺少解析库）。"""


def decode_text(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ImportError_("无法识别文本编码（支持 UTF-8 / GB18030 / UTF-16）")


def _csv_to_text(raw: bytes) -> str:
    text = decode_text(raw)
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return ""
    header = rows[0]
    lines = []
    for r in rows[1:]:
        cells = [f"{h}：{v}" for h, v in zip(header, r) if v and v.strip()]
        if cells:
            lines.append("；".join(cells) + "。")
    return "\n".join(lines)


def _docx_to_text(raw: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover - 依赖缺失分支
        raise ImportError_("解析 .docx 需要安装 python-docx：pip install python-docx") from e
    d = docx.Document(io.BytesIO(raw))
    parts = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            parts.append("#" * min(int(level), 6) + " " + t)
        else:
            parts.append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append("；".join(cells) + "。")
    return "\n\n".join(parts)


def _pdf_to_text(raw: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise ImportError_("解析 .pdf 需要安装 pypdf：pip install pypdf") from e
    reader = PdfReader(io.BytesIO(raw))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        t = (page.extract_text() or "").strip()
        if t:
            pages.append(f"# 第 {i} 页\n{t}")
    return "\n\n".join(pages)


def parse_upload(filename: str, raw: bytes) -> tuple[str, str]:
    """返回 (title, text)。抛 ImportError_ 说明失败原因。"""
    name = PurePosixPath((filename or "").replace("\\", "/")).name
    if not name:
        raise ImportError_("文件名为空")
    suffix = PurePosixPath(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ImportError_(f"不支持的文件类型 {suffix or '(无扩展名)'}，支持：{', '.join(SUPPORTED)}")
    if not raw:
        raise ImportError_(f"{name} 是空文件")
    if len(raw) > MAX_BYTES:
        raise ImportError_(f"{name} 超过 {MAX_BYTES // 1024 // 1024}MB 上限")
    if suffix in (".txt", ".md", ".markdown"):
        text = decode_text(raw)
    elif suffix == ".csv":
        text = _csv_to_text(raw)
    else:
        try:
            text = _docx_to_text(raw) if suffix == ".docx" else _pdf_to_text(raw)
        except ImportError_:
            raise  # 保留缺少可选解析依赖等明确提示。
        except Exception as e:
            # 第三方解析库使用不同异常类型；在文件解析边界统一为逐文件导入错误。
            raise ImportError_(f"{name} 解析失败：文件可能损坏、加密或格式不正确") from e
    text = text.replace("\r\n", "\n").strip()
    if len(text) < 20:
        raise ImportError_(f"{name} 解析后正文不足 20 字（扫描版 PDF 或空文档？）")
    title = PurePosixPath(name).stem
    # Markdown 首个一级标题优先作为文档标题
    for line in text.split("\n")[:5]:
        if line.startswith("# "):
            title = line[2:].strip() or title
            break
    return title, text
